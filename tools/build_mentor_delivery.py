"""Build mentor-facing Markdown, Word and PDF from the verified circular run.

Requires python-docx, reportlab and Pillow. Run from any directory.
Only presentation changes; numerical files and original figures are read-only.
"""
from pathlib import Path
import csv
import json
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
STEM = 'FAST_圆形MCF_佳伟老师交付版_20261003'
RUN = ROOT / 'runs/pilot/four_20261003_224047_034'
ROWS = list(csv.DictReader((RUN / 'metrics.csv').open()))
REPO = 'https://github.com/BoVVENGibNieAuf/LSA-FAST-Phase-Retrieval-Benchmark'


def metric(case, method, budget, key='field_nrmse'):
    found = [r for r in ROWS if r['case'] == case and r['method'] == method and int(r['budget']) == budget]
    assert len(found) == 1
    return float(found[0][key])


CASES = [('periodic_clean', '周期 / 无噪声'), ('aperiodic_clean', '非周期 / 无噪声'),
         ('periodic_shot_read', '周期 / 含噪声'), ('aperiodic_shot_read', '非周期 / 含噪声')]
main_table = [['布局 / 测量', '初始参考', 'HIO', 'ER', 'RAAR', 'L-BFGS']]
for case, label in CASES:
    main_table.append([label] + [f'{metric(case, m, 0 if m == "ZERO" else 200):.4f}'
                               for m in ['ZERO', 'HIO', 'ER', 'RAAR', 'LBFGS']])
noise_table = [['含噪声 ER', '80 次传播', '200 次传播']]
for case, label in CASES[2:]:
    noise_table.append([label.split(' / ')[0]] + [f'{metric(case, "ER", b):.4f}' for b in [80, 200]])

# Each item is a fixed page, shared by all three formats.
pages = [
    {'title': '圆形端面多芯光纤：四方法比较', 'landscape': False, 'items': [
        ('p', '阶段汇报｜2026年10月3日｜呈佳伟老师'),
        ('p', '佳伟老师：本阶段完成了圆形端面多芯光纤（MCF）的周期与非周期两类模拟，并在统一输入、初值和传播预算下比较 HIO、ER、RAAR 与 L-BFGS。16 项求解任务全部完成，几何检查和数值测试均通过。'),
        ('h', '主要发现'),
        ('p', '1. 无噪声时，两类布局均以 RAAR 的复场误差最低。'),
        ('p', '2. 含噪声时，四种算法中 ER 的复场误差最低；但从 80 次增加到 200 次传播后，误差反而上升。'),
        ('p', '3. 含噪声时，四种算法的相位误差均高于初始参考。当前结果提示，应继续检查噪声处理与停止准则。'),
        ('h', '核心结果：复场误差，越小越好'),
        ('table', main_table),
        ('small', '复场误差（NRMSE）衡量恢复的振幅与相位整体偏差，仅对齐一个全局相位。四算法均使用 200 次传播；“初始参考”为共同起点，未做迭代。'),
        ('h', '增加计算量未改善含噪声结果'),
        ('table', noise_table),
        ('p', '本轮为固定种子、单个混合样品的合成试运行，采用已知参考复场；上述比较限于本轮条件。'),
        ('p', '请佳伟老师指导下一阶段的工作重点。'),
    ]},
    {'title': '模型、比较条件与材料索引', 'landscape': False, 'items': [
        ('image', ROOT / 'runs/pilot/circular_20261003_224006_279/circular_geometry.png', 16.3, 10.9),
        ('small', '图1｜本次运行的实际布局。左：周期排列；右：非周期排列。黑实线为物理端面，红虚线为算法支持域（允许恢复光场的区域），蓝圈为纤芯。两组均为 163 芯，坐标单位为微米。'),
        ('h', '共同条件'),
        ('p', '物理端面半径 26 微米，芯中心范围 22 微米，算法支持域半径 29 微米。周期组芯间距 3.2 微米；非周期组最小芯间距 2.4 微米。端面外场能量检查为零。'),
        ('p', '两类布局均使用同一混合样品、已知合成参考校准和相同初始化规则。含噪声数据加入散粒噪声与读出噪声：总参考光子数 200000，读出噪声标准差 1 电子。'),
        ('p', '计算量按传播调用次数统一计数，包含 L-BFGS 线搜索中的拒绝试探；记录 80、200 次两个预算点。参数在评分前固定。四方法计算、评分及绘图共 32.62 秒，不含输入生成。'),
        ('h', '阅读与复跑'),
        ('p', '第3—6页为本次运行的原始恢复图。完整指标、方法卡和复跑说明保存在项目仓库；Word 可用于批注，PDF 用于固定版式阅读。'),
        ('link', '公开项目仓库（无需访问授权）', REPO),
        ('link', '完整结果包与校验信息', REPO + '/releases/tag/circular-results-2026-10-03'),
        ('small', '复跑入口：START_CIRCULAR_BENCHMARK.cmd。需要已激活的 MATLAB。本机运行与结果包恢复已验证；全新下载副本的 MATLAB 复跑尚未验证。完整方法卡见 docs/methods/FAST_四类求解器方法卡.docx。'),
    ]},
]

figure_notes = {
    'periodic_clean': '本组 RAAR 复场误差为 0.1182，为四算法最低；初始参考为 0.6555。',
    'aperiodic_clean': '本组 RAAR 复场误差为 0.1027，为四算法最低；初始参考为 0.6739。',
    'periodic_shot_read': '本组 ER 复场误差为 0.6313，为四算法最低；其相位误差为 0.4861 rad，高于初始参考的 0.2504 rad。',
    'aperiodic_shot_read': '本组 ER 复场误差为 0.6231，为四算法最低；其相位误差为 0.4705 rad，高于初始参考的 0.2584 rad。',
}
for i, (case, label) in enumerate(CASES, 2):
    pages.append({'title': label + '：200 次传播结果', 'landscape': True, 'items': [
        ('p', '上排为振幅，下排为参考校正相位。各列依次为仿真真值、初始参考、HIO、ER、RAAR、L-BFGS。'),
        ('image', RUN / case / 'fields_200.png', 26.6, 12.0),
        ('small', f'图{i}｜保留 MATLAB 原始图与色标。原图英文标题中的下划线被显示为下标，测量条件以本页中文标题为准。坐标单位：微米；相位单位：rad。'),
        ('p', figure_notes[case]),
    ]})


def build_word():
    doc = Document()
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(10.5 if name == 'Normal' else 19 if name == 'Title' else 12)
        style.paragraph_format.space_after = Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    for index, page in enumerate(pages):
        section = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        wide = page['landscape']
        section.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = (Cm(29.7), Cm(21)) if wide else (Cm(21), Cm(29.7))
        section.top_margin = section.bottom_margin = Cm(1.5)
        section.left_margin = section.right_margin = Cm(1.7)
        section.footer_distance = Cm(.7)
        if index == 0:
            foot = section.footer.paragraphs[0]
            foot.alignment = 2
            foot.add_run('圆形MCF阶段汇报  |  ')
            field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
        doc.add_heading(page['title'], 0 if index == 0 else 1)
        for item in page['items']:
            kind = item[0]
            if kind == 'table':
                t = doc.add_table(rows=0, cols=len(item[1][0])); t.style = 'Light Shading Accent 1'
                for values in item[1]:
                    row = t.add_row()
                    for cell, value in zip(row.cells, values):
                        cell.text = value
                        for run in cell.paragraphs[0].runs: run.font.size = Pt(9)
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            elif kind == 'image':
                with Image.open(item[1]) as im: w, h = im.size
                width = min(item[2], item[3] * w / h)
                doc.add_picture(str(item[1]), width=Cm(width))
            elif kind == 'h': doc.add_heading(item[1], 2)
            elif kind == 'link':
                p = doc.add_paragraph(item[1] + '：')
                relation = doc.part.relate_to(item[2], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), relation)
                run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '打开链接'; run.append(text); link.append(run); p._p.append(link)
            else:
                p = doc.add_paragraph(item[1])
                if kind == 'small':
                    for run in p.runs: run.font.size = Pt(9)
    doc.core_properties.title = '圆形端面多芯光纤四方法比较｜佳伟老师交付版'
    doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
    doc.save(OUT / (STEM + '.docx'))


def build_pdf():
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    c = canvas.Canvas(str(OUT / (STEM + '.pdf')))
    c.setTitle('圆形端面多芯光纤四方法比较｜佳伟老师交付版'); c.setAuthor('')
    for index, page in enumerate(pages):
        width, height = landscape(A4) if page['landscape'] else A4
        c.setPageSize((width, height)); x = 48; y = height - 44; available = width - 96
        def para(text, size=10.5, leading=16, gap=7, color='#25334A'):
            nonlocal y
            style = ParagraphStyle('zh', fontName='STSong-Light', fontSize=size, leading=leading,
                                   wordWrap='CJK', textColor=colors.HexColor(color))
            p = Paragraph(text, style); _, h = p.wrap(available, height)
            assert y - h > 38, (index, text[:50], y, h)
            p.drawOn(c, x, y - h); y -= h + gap
        para(escape(page['title']), 19, 25, 13, '#163D68')
        for item in page['items']:
            kind = item[0]
            if kind == 'table':
                data = item[1]; n = len(data[0]); col_widths = [available / n] * n
                if n == 6: col_widths = [114] + [(available - 114) / 5] * 5
                t = Table(data, colWidths=col_widths)
                t.setStyle(TableStyle([('FONTNAME', (0,0),(-1,-1),'STSong-Light'),
                    ('FONTSIZE',(0,0),(-1,-1),9), ('LEADING',(0,0),(-1,-1),12),
                    ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E5EDF5')),
                    ('TEXTCOLOR',(0,0),(-1,-1),colors.HexColor('#25334A')),
                    ('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),
                    ('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#CAD5E0'))]))
                _, h = t.wrap(available, height); assert y - h > 38
                t.drawOn(c,x,y-h); y -= h + 9
            elif kind == 'image':
                with Image.open(item[1]) as im: w,h = im.size
                scale = min(available/w, item[2]*28.346/w, item[3]*28.346/h)
                dw,dh = w*scale,h*scale; assert y-dh > 38
                c.drawImage(str(item[1]),x+(available-dw)/2,y-dh,width=dw,height=dh); y-=dh+8
            elif kind == 'h': para(escape(item[1]),12,18,7,'#163D68')
            elif kind == 'small': para(escape(item[1]),9,13,8)
            elif kind == 'link': para(f'<link href="{escape(item[2])}" color="#145B9A">{escape(item[1])}</link>',10,15,5)
            else: para(escape(item[1]))
        c.setFont('STSong-Light',8); c.setFillColor(colors.HexColor('#68788A'))
        c.drawString(48,23,'圆形MCF阶段汇报 | 2026-10-03'); c.drawRightString(width-48,23,f'{index+1} / {len(pages)}')
        c.showPage()
    c.save()


def build_markdown():
    out = []
    for page in pages:
        out += ['# ' + page['title'], '']
        for item in page['items']:
            kind = item[0]
            if kind == 'table':
                data = item[1]; out += ['|'+'|'.join(data[0])+'|', '|'+'|'.join(['---']*len(data[0]))+'|']
                out += ['|'+'|'.join(row)+'|' for row in data[1:]]
            elif kind == 'image': out.append(f'![实际运行图](../{item[1].relative_to(ROOT).as_posix()})')
            elif kind == 'h': out.append('## ' + item[1])
            elif kind == 'link': out.append(f'[{item[1]}]({item[2]})')
            else: out.append(item[1])
            out.append('')
    (OUT / (STEM + '.md')).write_text('\n'.join(out))


if __name__ == '__main__':
    build_word(); build_pdf(); build_markdown()
    print(json.dumps({'stem': STEM, 'pages': len(pages), 'metric_rows': len(ROWS)}, ensure_ascii=False))
